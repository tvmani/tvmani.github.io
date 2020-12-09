from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import load_workbook
import collections as col
import string
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import datetime
import os
from os.path import basename, splitext
from string import Template

# todo: filename should be passed in commandline for function
currentDirectory = "E:\\2020\\git\\quiz_app\\quizapp\\src\\tools"
filename = "12_TM_Botany_U01.xlsx"
# lessonTitle = "அலகு 5 - தாவர செயலியல் (Plant Physiology)"
lessonTitle = "தாவரங்களில் இனப்பெருக்கம்"
wb = load_workbook(currentDirectory + "\\" + filename)

sheet_ranges = wb.worksheets[0]

print(sheet_ranges.tables.items())

(table_name, table_cell_range) = sheet_ranges.tables.items()[0]

print(table_name, table_cell_range, sep=",")


def getTotalRecords(tableRange):
    # A1:C5
    (lower_cell, upper_cell) = tableRange.split(":")
    start_row = int(lower_cell[1:])
    end_row = int(upper_cell[1:])
    return end_row - start_row


def get_cell_identifier(tableRange):
    # A1:C5
    (lower_cell, upper_cell) = tableRange.split(":")
    start_row_id = lower_cell[0]
    end_row_id = upper_cell[0]
    return (start_row_id, end_row_id)


def get_cell_indexes(tableRange):
    # A1:C5
    (lower_cell, upper_cell) = tableRange.split(":")
    start_row_index = int(lower_cell[1:])
    end_row_index = int(upper_cell[1:]) - 1
    return (start_row_index, end_row_index)


def getOptionAsString(optionInput):
    result = ""
    if optionInput is None:
        result = ""
    else:
        result = str(optionInput)
    return result


def getQuizAsJSObject(quizRecord):
    print(quizRecord.question)
    quizAsJSObject = """{{
title: "{question}",
subtitle: "{question_no}",
choices: [
    `{option_1}`,
    `{option_2}`,
    `{option_3}`,
    `{option_4}`,
    `{option_5}`,
],
solution: `{correct_answer}`,
explanation: `{topic_name}`
}}
""".format(question=quizRecord.question,
           question_no=quizRecord.question_no,
           option_1=quizRecord.option_1,
           option_2=quizRecord.option_2,
           option_3=quizRecord.option_3,
           option_4=quizRecord.option_4,
           option_5=quizRecord.option_5,
           correct_answer=quizRecord.correct_answer,
           topic_name=quizRecord.topic_name,
           )
    return quizAsJSObject


def createQuizFile(allQuizRecords):
    quizContentHeader = """export default {
	title: "Week 4 Quiz",
	category: "12th Biology Botany",
	challenges: [
    """

    quizContentFooter = """]
    };
    """

    quizAsJSFileContent = quizContentHeader + \
        allQuizRecords + \
        quizContentFooter

    return quizAsJSFileContent.encode('utf8')


def getQuizFilename(csvFilename):
    filenameWithoutExtn = splitext(basename(csvFilename))[0]
    return filenameWithoutExtn + "." + "js"


def writeToFile(allQuizRecords, filename):
    f = open(filename, "wb")
    f.write(createQuizFile(allQuizRecords))
    f.close()


print(getTotalRecords(table_cell_range))
print(get_cell_identifier(table_cell_range))
print(get_cell_indexes(table_cell_range))
excel_A_to_Z = list(string.ascii_uppercase)

(start_row_index, end_row_index) = get_cell_indexes(table_cell_range)
(start_row_id, end_row_id) = get_cell_identifier(table_cell_range)

QuizRecord = col.namedtuple('QuizRecord', 'question_no, question, exams_already_asked, year, option_1, \
                            option_2, option_3, option_4, option_5, correct_answer, page_no,topic_name')
QuizRecordsDb = []

for row in range(start_row_index, end_row_index + 1, 1):
    print("processing row " + str(start_row_index))
    quizRecord = QuizRecord(sheet_ranges['A' + str(row + 1)].value,
                            sheet_ranges['B' + str(row + 1)].value,
                            sheet_ranges['C' + str(row + 1)].value,
                            sheet_ranges['D' + str(row + 1)].value,
                            sheet_ranges['E' + str(row + 1)].value,
                            sheet_ranges['F' + str(row + 1)].value,
                            sheet_ranges['G' + str(row + 1)].value,
                            sheet_ranges['H' + str(row + 1)].value,
                            sheet_ranges['I' + str(row + 1)].value,
                            sheet_ranges['J' + str(row + 1)].value,
                            sheet_ranges['K' + str(row + 1)].value,
                            sheet_ranges['L' + str(row + 1)].value,
                            )
    QuizRecordsDb.append(quizRecord)

print(len(QuizRecordsDb))

quizFilename = getQuizFilename(filename)
allQuizRecords = ""
quiz_separator = ","
for quizRecord in QuizRecordsDb:
    # print(quizRecord)
    #writeToFile(quizRecord, quizFilename)
    if len(allQuizRecords) == 0:
        quiz_separator = ""
    else:
        quiz_separator = ","
    allQuizRecords = allQuizRecords + quiz_separator + getQuizAsJSObject(quizRecord)


writeToFile(allQuizRecords, quizFilename)
print("Completed creation of quiz js")
