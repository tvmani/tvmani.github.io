from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import load_workbook

wb = load_workbook(filename = r"E:\Resources\RangeBook.xlsx")

# my_range = wb.defined_names['Fruits_Price']

sheetOne = wb['One']

print(sheetOne.tables.items())

tblFruitsPrice = sheetOne.tables.get("Fruits_Price")

print(tblFruitsPrice)

# for table in sheetOne.tables.values():
# #     print(table)
# for column in table.:
#     column.name = value